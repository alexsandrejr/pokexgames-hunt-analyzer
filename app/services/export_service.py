"""Exportação de tabelas para CSV, Excel (.xlsx), PDF e JSON.

Um ``ExportDocument`` reúne uma ou mais ``ExportTable``. Os valores das linhas
são **brutos** (números, datas, segundos); cada coluna informa seu ``ValueKind``,
que decide a formatação em cada formato:

* CSV: separador e decimal do formato de números ativo — pt-BR: ``;`` e vírgula
  (abre direto no Excel em português); en-US: ``,`` e ponto;
* Excel: números reais com formatação de milhar, datas e durações ``[h]:mm:ss``;
* PDF: valores formatados como na interface;
* JSON: valores brutos (durações em segundos, datas em ISO 8601).
"""

from __future__ import annotations

import csv
import html
import json
import re
from dataclasses import dataclass, field
from datetime import datetime
from enum import Enum
from pathlib import Path
from typing import Any

from app.utils.constants import APP_NAME
from app.utils.formatters import (
    ValueKind,
    format_datetime,
    format_duration,
    format_value,
    number_format,
)

DECIMALS = 2


class ExportFormat(Enum):
    CSV = ("CSV", "csv", "CSV")
    XLSX = ("Excel", "xlsx", "Planilha do Excel")
    PDF = ("PDF", "pdf", "Documento PDF")
    JSON = ("JSON", "json", "JSON")

    def __init__(self, label: str, extension: str, description: str) -> None:
        self.label = label
        self.extension = extension
        self.description = description

    @property
    def file_filter(self) -> str:
        return f"{self.description} (*.{self.extension})"


@dataclass(frozen=True)
class ExportColumn:
    header: str
    # ``None``: o tipo vem da linha (``ExportTable.row_kinds``), em tabelas transpostas.
    kind: ValueKind | None = ValueKind.TEXT


@dataclass
class ExportTable:
    title: str
    columns: list[ExportColumn]
    rows: list[list[Any]] = field(default_factory=list)
    # Tabelas "transpostas" (ex.: comparação): tipo de cada linha, usado nas colunas sem tipo.
    row_kinds: list[ValueKind] | None = None

    def cell_kind(self, row_index: int, column_index: int) -> ValueKind:
        kind = self.columns[column_index].kind
        if kind is None:
            return self.row_kinds[row_index] if self.row_kinds else ValueKind.TEXT
        return kind

    def display_rows(self) -> list[list[str]]:
        return [[format_value(self.cell_kind(r, c), value) for c, value in enumerate(row)]
                for r, row in enumerate(self.rows)]


@dataclass
class ExportDocument:
    title: str
    tables: list[ExportTable]
    subtitle: str = ""
    generated_at: datetime = field(default_factory=datetime.now)


class ExportError(Exception):
    """Falha ao gravar o arquivo, com mensagem pronta para o usuário."""


def export_document(document: ExportDocument, path: str | Path, fmt: ExportFormat) -> Path:
    target = Path(path)
    if target.suffix.lower() != f".{fmt.extension}":
        target = target.with_name(f"{target.name}.{fmt.extension}")
    writers = {ExportFormat.CSV: _write_csv, ExportFormat.XLSX: _write_xlsx,
               ExportFormat.PDF: _write_pdf, ExportFormat.JSON: _write_json}
    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        writers[fmt](document, target)
    except PermissionError as exc:
        raise ExportError(f"Sem permissão para gravar {target}. O arquivo está aberto "
                          "em outro programa?") from exc
    except OSError as exc:
        raise ExportError(f"Não foi possível gravar {target}: {exc}") from exc
    return target


def suggested_filename(title: str, fmt: ExportFormat, moment: datetime | None = None) -> str:
    slug = re.sub(r"[^\w]+", "_", title.lower(), flags=re.UNICODE).strip("_") or "export"
    stamp = (moment or datetime.now()).strftime("%Y-%m-%d_%H-%M")
    return f"{slug}_{stamp}.{fmt.extension}"


# ----------------------------------------------------------------------- CSV

def _csv_value(kind: ValueKind, value: Any) -> str:
    if value is None:
        return ""
    if kind is ValueKind.DURATION:
        return format_duration(value)
    if isinstance(value, datetime):
        return format_datetime(value)
    if isinstance(value, bool):
        return "Sim" if value else "Não"
    if isinstance(value, float):
        text = f"{round(value, DECIMALS):.{DECIMALS}f}".rstrip("0").rstrip(".")
        return text.replace(".", number_format().decimal)
    return str(value)


def _write_csv(document: ExportDocument, path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        # Com vírgula decimal, o separador de campos precisa ser ";".
        delimiter = ";" if number_format().decimal == "," else ","
        writer = csv.writer(file, delimiter=delimiter)
        several = len(document.tables) > 1
        for index, table in enumerate(document.tables):
            if several:
                if index:
                    writer.writerow([])
                writer.writerow([table.title])
            writer.writerow([column.header for column in table.columns])
            for r, row in enumerate(table.rows):
                writer.writerow([_csv_value(table.cell_kind(r, c), value)
                                 for c, value in enumerate(row)])


# ---------------------------------------------------------------------- JSON

def _json_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, float):
        return round(value, DECIMALS)
    return value


def _json_key(column: ExportColumn) -> str:
    return f"{column.header} (s)" if column.kind is ValueKind.DURATION else column.header


def _write_json(document: ExportDocument, path: Path) -> None:
    payload = {
        "title": document.title,
        "subtitle": document.subtitle,
        "generated_at": document.generated_at.isoformat(timespec="seconds"),
        "generator": APP_NAME,
        "tables": [
            {
                "title": table.title,
                "rows": [
                    {_json_key(column): _json_value(value)
                     for column, value in zip(table.columns, row, strict=False)}
                    for row in table.rows
                ],
            }
            for table in document.tables
        ],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


# --------------------------------------------------------------------- Excel

_XLSX_FORMATS = {
    ValueKind.MONEY: "#,##0",
    ValueKind.NUMBER: "#,##0",
    ValueKind.DURATION: "[h]:mm:ss",
    ValueKind.DATETIME: "dd/mm/yyyy hh:mm",
    ValueKind.PERCENT: "0.0%",
}


def _sheet_title(title: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", title).strip()[:31] or "Tabela"
    name, counter = base, 2
    while name in used:
        suffix = f" ({counter})"
        name, counter = base[: 31 - len(suffix)] + suffix, counter + 1
    used.add(name)
    return name


def _write_xlsx(document: ExportDocument, path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:  # pragma: no cover - depende do ambiente
        raise ExportError("A exportação para Excel requer o pacote openpyxl "
                          "(pip install -r requirements.txt).") from exc

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = document.title
    workbook.properties.creator = APP_NAME
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2A3140")
    used_titles: set[str] = set()

    for table in document.tables:
        sheet = workbook.create_sheet(_sheet_title(table.title, used_titles))
        sheet.append([column.header for column in table.columns])
        for cell in sheet[1]:
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(vertical="center")
        displays = table.display_rows()
        widths = [len(column.header) for column in table.columns]
        for r, row in enumerate(table.rows):
            values = []
            for c, value in enumerate(row):
                kind = table.cell_kind(r, c)
                if kind is ValueKind.DURATION and value is not None:
                    value = value / 86400  # Excel guarda tempo como fração de dia
                values.append(value)
                widths[c] = max(widths[c], len(displays[r][c]))
            sheet.append(values)
            for c, cell in enumerate(sheet[sheet.max_row]):
                number_format = _XLSX_FORMATS.get(table.cell_kind(r, c))
                if number_format:
                    cell.number_format = number_format
        for c, width in enumerate(widths, start=1):
            sheet.column_dimensions[sheet.cell(1, c).column_letter].width = min(50, width + 3)
        sheet.freeze_panes = "A2"
    workbook.save(path)


# ----------------------------------------------------------------------- PDF

_NUMERIC_KINDS = {ValueKind.NUMBER, ValueKind.MONEY, ValueKind.DURATION, ValueKind.PERCENT}


def document_html(document: ExportDocument) -> str:
    """HTML simples (tema claro, próprio para impressão) usado no PDF."""
    parts = [
        "<html><body style='font-family: Segoe UI, Arial; color: #111827;'>",
        f"<h2 style='margin-bottom: 0'>{html.escape(document.title)}</h2>",
        "<p style='color: #6b7280; margin-top: 2px'>"
        f"{html.escape(document.subtitle)}{' · ' if document.subtitle else ''}"
        f"Gerado em {format_datetime(document.generated_at)} · {html.escape(APP_NAME)}</p>",
    ]
    for table in document.tables:
        parts.append(f"<h3 style='margin-top: 14px'>{html.escape(table.title)}</h3>")
        parts.append("<table width='100%' cellspacing='0' cellpadding='4' "
                     "style='border-collapse: collapse; font-size: 8pt;'>")
        parts.append("<tr style='background-color: #e5e7eb;'>")
        for column in table.columns:
            numeric = column.kind is None or column.kind in _NUMERIC_KINDS
            align = "right" if numeric else "left"
            parts.append(f"<th align='{align}'>{html.escape(column.header)}</th>")
        parts.append("</tr>")
        for r, row in enumerate(table.display_rows()):
            shade = " style='background-color: #f5f6f8;'" if r % 2 else ""
            parts.append(f"<tr{shade}>")
            for c, text in enumerate(row):
                transposed_value = table.columns[c].kind is None
                numeric = table.cell_kind(r, c) in _NUMERIC_KINDS
                align = "right" if transposed_value or numeric else "left"
                parts.append(f"<td align='{align}'>{html.escape(text)}</td>")
            parts.append("</tr>")
        if not table.rows:
            parts.append(f"<tr><td colspan='{len(table.columns)}' style='color: #6b7280;'>"
                         "Sem registros.</td></tr>")
        parts.append("</table>")
    parts.append("</body></html>")
    return "".join(parts)


def _write_pdf(document: ExportDocument, path: Path) -> None:
    # Qt só é importado aqui: o restante do módulo não depende de interface.
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QPageLayout, QPageSize, QPdfWriter, QTextDocument

    widest = max((len(table.columns) for table in document.tables), default=0)
    orientation = (QPageLayout.Orientation.Landscape if widest > 7
                   else QPageLayout.Orientation.Portrait)
    writer = QPdfWriter(str(path))
    writer.setTitle(document.title)
    writer.setCreator(APP_NAME)
    writer.setPageLayout(QPageLayout(QPageSize(QPageSize.PageSizeId.A4), orientation,
                                     QMarginsF(12, 12, 12, 12), QPageLayout.Unit.Millimeter))
    text_document = QTextDocument()
    text_document.setHtml(document_html(document))
    text_document.print_(writer)
    if not path.exists() or path.stat().st_size == 0:
        raise OSError("o PDF não foi gerado")

