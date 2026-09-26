"""Interface da Fase 3: comparação, relatórios de item/inimigo, pesquisa avançada e exportação."""

import csv
import os
import tempfile
import unittest
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook  # noqa: E402
from PySide6.QtCore import QItemSelectionModel  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

from app.services.entity_report import EntityKind  # noqa: E402
from app.services.export_documents import COMPARISON_ROWS  # noqa: E402
from app.services.export_service import ExportFormat  # noqa: E402
from app.services.filters import HuntFilter, ItemCondition, Operator  # noqa: E402
from app.services.hunt_service import HuntService  # noqa: E402
from app.services.statistics_service import StatisticsService  # noqa: E402
from app.ui.main_window import MainWindow  # noqa: E402
from app.ui.theme import apply_theme  # noqa: E402
from tests.helpers import memory_database  # noqa: E402
from tests.test_filters import make_hunt  # noqa: E402


class Phase3UiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])
        apply_theme(cls.app)

    def setUp(self) -> None:
        self.database = memory_database()
        self.service = HuntService(self.database)
        for text in (
            make_hunt(1, "2026-09-01 10:00:00", kills=250, profit=800_000,
                      drops={"nightmare ore": 40, "nightmare gem": 1200},
                      enemies={"Mecha Charizard": 30}),
            make_hunt(2, "2026-09-10 10:00:00", duration=7200, kills=640, profit=2_000_000,
                      drops={"nightmare ore": 80, "nightmare gem": 900},
                      enemies={"Mecha Charizard": 50}),
            make_hunt(3, "2026-09-11 10:00:00", duration=1800, kills=160, profit=-50_000,
                      drops={"metal scraps": 5}, enemies={"Mecha Venusaur": 20},
                      supplies={"Healing Elixir": 4}),
        ):
            self.service.import_text(text)
        self.window = MainWindow(self.service, StatisticsService(self.database), Path("m.db"))
        self.window.show()
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self) -> None:
        self.window.close()
        self.window.deleteLater()
        self.database.dispose()
        self._tmp.cleanup()

    def select_rows(self, *rows: int) -> None:
        page = self.window.hunts_page
        selection = page._table.selectionModel()
        selection.clearSelection()
        for row in rows:
            selection.select(page._table.proxy.index(row, 0),
                             QItemSelectionModel.SelectionFlag.Select
                             | QItemSelectionModel.SelectionFlag.Rows)

    # --------------------------------------------------------------- comparação

    def test_compare_selected_hunts(self) -> None:
        self.window.show_page("hunts")
        page = self.window.hunts_page
        self.select_rows(0)
        self.assertFalse(page._compare_button.isEnabled())
        self.select_rows(0, 2)
        self.assertTrue(page._compare_button.isEnabled())

        dialog = page.compare_selected()
        table = dialog.table
        self.assertEqual(table.rowCount(), len(COMPARISON_ROWS))
        self.assertEqual(table.columnCount(), 3)
        self.assertEqual([h.session_id for h in dialog.hunts], [1, 3])  # cronológico
        labels = [table.item(r, 0).text() for r in range(table.rowCount())]
        profit_row = labels.index("Profit")
        self.assertEqual(table.item(profit_row, 1).text(), "800.000")
        self.assertEqual(table.item(profit_row, 2).text(), "-50.000")
        self.assertEqual(table.item(labels.index("Duração"), 2).text(), "00:30:00")

        path = dialog.export_button.export_to(
            dialog.export_button._provider(), self.dir / "cmp", ExportFormat.XLSX)
        sheet = load_workbook(path).active
        self.assertEqual(sheet.cell(1, 1).value, "Métrica")
        dialog.close()

    # --------------------------------------------------------------- exportação

    def test_hunts_export_follows_search_and_order(self) -> None:
        self.window.show_page("hunts")
        page = self.window.hunts_page
        path = page.export_button.export_to(page._export_document(), self.dir / "h.csv",
                                            ExportFormat.CSV)
        with path.open(encoding="utf-8-sig", newline="") as file:
            rows = list(csv.reader(file, delimiter=";"))
        self.assertEqual([r[1] for r in rows[1:]], ["3", "2", "1"])  # mais recentes primeiro
        self.assertIn("Exportado:", self.window.statusBar().currentMessage())

        page._search.setText("10/09")
        path = page.export_button.export_to(page._export_document(), self.dir / "s.csv",
                                            ExportFormat.CSV)
        with path.open(encoding="utf-8-sig", newline="") as file:
            self.assertEqual(len(list(csv.reader(file, delimiter=";"))), 2)

    def test_reports_export(self) -> None:
        self.window.filters.set_filter(HuntFilter(items=(ItemCondition("nightmare ore"),)))
        self.window.show_page("reports")
        page = self.window.reports_page
        document = page._export_document()
        self.assertIn("Drop: nightmare ore", document.subtitle)
        path = page.export_button.export_to(document, self.dir / "r.pdf", ExportFormat.PDF)
        self.assertTrue(path.read_bytes().startswith(b"%PDF"))

    # ------------------------------------------------------- item e inimigo

    def test_item_report_page(self) -> None:
        self.window.show_page("items")
        page = self.window.pages["items"]
        self.assertFalse(page.export_button.isEnabled())
        names = [page.name_combo.itemText(i) for i in range(page.name_combo.count())]
        self.assertIn("nightmare ore", names)

        page.analyze("nightmare ore")
        self.assertEqual(page._stack.currentIndex(), 1)
        self.assertEqual(page._cards["Quantidade total"]._value.text(), "120")
        self.assertEqual(page._cards["Hunts"]._value.text(), "2 de 3")
        self.assertEqual(page._cards["Média por Hunt"]._value.text(), "60")
        self.assertEqual(page._cards["Média por hora"]._value.text(), "40")  # 120 em 3h
        self.assertEqual(page._cards["Maior quantidade"]._value.text(), "80")
        self.assertEqual(page._cards["Valor total"]._value.text(), "12.000")
        self.assertEqual([p.value for p in page._chart._points], [40, 80, 0])
        self.assertEqual(page._table.source_model.rowCount(), 2)
        self.assertTrue(page.export_button.isEnabled())

        # O filtro global é respeitado ao atualizar.
        self.window.filters.set_filter(HuntFilter(player="nobody-here"))
        self.assertEqual(page._stack.currentIndex(), 0)
        self.window.filters.clear()
        self.assertEqual(page._stack.currentIndex(), 1)

    def test_supply_and_unknown_item(self) -> None:
        self.window.show_page("items")
        page = self.window.pages["items"]
        page.kind_combo.setCurrentIndex(page.kind_combo.findData(EntityKind.SUPPLIES))
        page.analyze("Healing Elixir")
        self.assertEqual(page._cards["Quantidade total"]._value.text(), "4")
        page.analyze("não existe")
        self.assertEqual(page._stack.currentIndex(), 0)
        self.assertEqual(page._empty._heading.text(), "Nada encontrado")
        self.assertFalse(page.export_button.isEnabled())

    def test_enemy_report_page(self) -> None:
        self.window.show_page("enemies")
        page = self.window.pages["enemies"]
        self.assertFalse(page.kind_combo.isVisibleTo(page))
        page.analyze("mecha charizard")
        self.assertEqual(page._cards["Total derrotado"]._value.text(), "80")
        self.assertNotIn("Valor total", page._cards)
        path = page.export_button.export_to(page._export_document(), self.dir / "e.json",
                                            ExportFormat.JSON)
        self.assertIn("Inimigo: mecha charizard", path.read_text(encoding="utf-8"))

    # ------------------------------------------------------- pesquisa avançada

    def test_advanced_search_in_panel(self) -> None:
        self.window.show_page("hunts")
        panel = self.window.filter_panel
        panel._advanced.setText("nightmare ore > 50 e nightmare gem > 800")
        self.assertTrue(panel.apply())
        self.assertEqual(self.window.filters.current.items, (
            ItemCondition("nightmare ore", Operator.GT, 50),
            ItemCondition("nightmare gem", Operator.GT, 800),
        ))
        self.assertEqual(self.window.hunts_page._table.visible_row_count(), 1)
        self.assertIn("Drop: nightmare ore > 50", panel._summary.full_text())

        panel._advanced.setText("coisa inexistente > 3")
        self.assertFalse(panel.apply())
        self.assertTrue(panel._advanced.property("invalid"))
        self.assertIn("não foi encontrado", panel._error.text())

        panel.clear()
        self.assertEqual(panel._advanced.text(), "")
        self.assertTrue(self.window.filters.current.is_empty)


if __name__ == "__main__":
    unittest.main()
