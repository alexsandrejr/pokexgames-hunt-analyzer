"""Exportação: cada formato é gravado e lido de volta."""

import csv
import json
import os
import tempfile
import unittest
from datetime import datetime, timedelta
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook  # noqa: E402
from PySide6.QtWidgets import QApplication  # noqa: E402

from app.services.entity_report import EntityKind  # noqa: E402
from app.services.export_documents import (  # noqa: E402
    comparison_document,
    comparison_table,
    entity_document,
    hunts_document,
    report_document,
)
from app.services.export_service import (  # noqa: E402
    ExportError,
    ExportFormat,
    export_document,
    suggested_filename,
)
from app.services.hunt_service import HuntService  # noqa: E402
from app.services.statistics_service import StatisticsService  # noqa: E402
from tests.helpers import SAMPLE_HUNT_PATH, memory_database, sample_text  # noqa: E402


class ExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.app = QApplication.instance() or QApplication([])  # necessário para o PDF
        cls.database = memory_database()
        cls.service = HuntService(cls.database)
        cls.stats = StatisticsService(cls.database)
        cls.service.import_file(SAMPLE_HUNT_PATH)
        cls.service.import_text(sample_text(**{"Session ID": 1081, "Start": "2026-09-27 10:00:00",
                                              "Duration seconds": 90061, "Profit": -5000}))
        cls.hunts = cls.service.get_summaries(h.id for h in cls.service.list_hunts())

    @classmethod
    def tearDownClass(cls) -> None:
        cls.database.dispose()

    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_csv(self) -> None:
        path = export_document(hunts_document(self.hunts, ["Player: Royal"]),
                               self.dir / "hunts", ExportFormat.CSV)
        self.assertEqual(path.suffix, ".csv")
        with path.open(encoding="utf-8-sig", newline="") as file:
            rows = list(csv.reader(file, delimiter=";"))
        header, first, second = rows
        record = dict(zip(header, first, strict=True))
        self.assertEqual(record["Data"], "26/09/2026 15:24:20")
        self.assertEqual(record["Duração"], "01:11:50")
        self.assertEqual(record["Profit"], "1119522")
        self.assertEqual(record["Profit/h"], "935099")  # inteiro no JSON de origem
        self.assertEqual(record["Arquivo"], "sample_hunt.json")
        self.assertEqual(dict(zip(header, second, strict=True))["Duração"], "25:01:01")

    def test_csv_decimal_comma_and_sections(self) -> None:
        document = report_document(self.stats.overview(), self.stats.metric_breakdown(),
                                   self.stats.top_drops(), self.stats.top_supplies(),
                                   self.stats.top_enemies())
        path = export_document(document, self.dir / "r.csv", ExportFormat.CSV)
        text = path.read_text(encoding="utf-8-sig")
        self.assertTrue(text.startswith("Resumo\n"))  # read_text converte \r\n em \n
        self.assertIn("\n\nMétricas\n", text)
        self.assertIn("\n\nInimigos\n", text)
        self.assertRegex(text, r"Profit;\d+;\d+(,\d+)?;")  # decimais com vírgula

    def test_json(self) -> None:
        path = export_document(hunts_document(self.hunts), self.dir / "h.json", ExportFormat.JSON)
        payload = json.loads(path.read_text(encoding="utf-8"))
        self.assertEqual(payload["title"], "Hunts")
        row = payload["tables"][0]["rows"][0]
        self.assertEqual(row["Profit"], 1119522)
        self.assertEqual(row["Duração (s)"], 4310)
        self.assertEqual(row["Data"], "2026-09-26T15:24:20")
        self.assertEqual(payload["tables"][0]["rows"][1]["Profit"], -5000)

    def test_xlsx(self) -> None:
        document = report_document(self.stats.overview(), self.stats.metric_breakdown(),
                                   self.stats.top_drops(), self.stats.top_supplies(),
                                   self.stats.top_enemies())
        document.tables.insert(0, hunts_document(self.hunts).tables[0])
        path = export_document(document, self.dir / "r.xlsx", ExportFormat.XLSX)
        workbook = load_workbook(path)
        self.assertEqual(workbook.sheetnames,
                         ["Hunts", "Resumo", "Métricas", "Drops", "Supplies", "Inimigos"])
        sheet = workbook["Hunts"]
        headers = [cell.value for cell in sheet[1]]
        profit = sheet.cell(2, headers.index("Profit") + 1)
        self.assertEqual(profit.value, 1119522)
        self.assertEqual(profit.number_format, "#,##0")
        duration = sheet.cell(2, headers.index("Duração") + 1)
        self.assertEqual(duration.value, timedelta(seconds=4310))  # duração real no Excel
        self.assertEqual(duration.number_format, "[h]:mm:ss")
        self.assertEqual(sheet.cell(2, 1).value, datetime(2026, 9, 26, 15, 24, 20))
        self.assertEqual(sheet.freeze_panes, "A2")
        metrics = workbook["Métricas"]
        self.assertEqual(metrics.cell(2, 1).value, "Duração")
        self.assertEqual(metrics.cell(2, 2).number_format, "[h]:mm:ss")  # tipo vem da linha
        self.assertEqual(metrics.cell(4, 8).value, "/h")  # coluna de texto não herda o tipo

    def test_pdf(self) -> None:
        document = comparison_document(self.hunts)
        path = export_document(document, self.dir / "c.pdf", ExportFormat.PDF)
        data = path.read_bytes()
        self.assertTrue(data.startswith(b"%PDF"))
        self.assertGreater(len(data), 1000)

    def test_comparison_table(self) -> None:
        table = comparison_table(self.hunts)
        self.assertEqual([c.header for c in table.columns][0], "Métrica")
        self.assertTrue(table.columns[1].header.startswith("Hunt 1080 · 26/09/2026"))
        display = {row[0]: row[1:] for row in table.display_rows()}
        self.assertEqual(display["Duração"], ["01:11:50", "25:01:01"])
        self.assertEqual(display["Profit"], ["1.119.522", "-5.000"])
        self.assertEqual(display["Player"], ["Royalzxd", "Royalzxd"])

    def test_entity_document(self) -> None:
        report = self.stats.entity_report(EntityKind.DROPS, "metal scraps")
        document = entity_document(report, ["Player: Royal"])
        self.assertEqual(document.title, "Drop: metal scraps")
        summary = {row[0]: row[1] for row in document.tables[0].display_rows()}
        self.assertEqual(summary["Quantidade total"], "388")
        self.assertEqual(summary["Valor total"], "776.000")
        self.assertEqual(summary["Presença"], "100,0%")
        enemy = entity_document(self.stats.entity_report(EntityKind.ENEMIES, "Mecha Charizard"))
        self.assertNotIn("Valor", [c.header for c in enemy.tables[1].columns])
        for fmt in ExportFormat:
            with self.subTest(fmt=fmt):
                self.assertTrue(export_document(document, self.dir / f"e.{fmt.extension}",
                                                fmt).exists())

    def test_empty_tables_export(self) -> None:
        document = hunts_document([])
        for fmt in ExportFormat:
            with self.subTest(fmt=fmt):
                self.assertTrue(export_document(document, self.dir / f"v.{fmt.extension}",
                                                fmt).exists())

    def test_write_error_is_reported(self) -> None:
        blocker = self.dir / "arquivo"
        blocker.write_text("x")
        with self.assertRaises(ExportError):
            export_document(hunts_document(self.hunts), blocker / "sub" / "h.csv",
                            ExportFormat.CSV)

    def test_suggested_filename(self) -> None:
        moment = datetime(2026, 9, 26, 15, 30)
        self.assertEqual(suggested_filename("Relatório de Hunts", ExportFormat.XLSX, moment),
                         "relatório_de_hunts_2026-09-26_15-30.xlsx")
        self.assertEqual(suggested_filename("Drop: nightmare ore", ExportFormat.CSV, moment),
                         "drop_nightmare_ore_2026-09-26_15-30.csv")


if __name__ == "__main__":
    unittest.main()
